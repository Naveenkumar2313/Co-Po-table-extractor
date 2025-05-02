from flask import Flask, request, send_file, render_template
import os
from werkzeug.utils import secure_filename
from docx import Document
from openpyxl import Workbook

app = Flask(__name__)
UPLOAD_FOLDER = 'uploads'
OUTPUT_FOLDER = 'output'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return 'No file part'

    file = request.files['file']
    if file.filename == '':
        return 'No selected file'

    filename = secure_filename(file.filename)
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    file.save(filepath)

    # Load the Word document
    doc = Document(filepath)
    tables = doc.tables

    if not tables:
        return 'No tables found in document.'

    last_table = tables[-1]

    wb = Workbook()
    ws = wb.active
    ws.title = "CO-PO Mapping"

    for row in last_table.rows:
        row_data = [cell.text.strip() for cell in row.cells]
        ws.append(row_data)

    output_path = os.path.join(OUTPUT_FOLDER, 'co-po.xlsx')
    wb.save(output_path)

    return send_file(output_path, as_attachment=True)

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
